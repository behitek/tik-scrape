import logging
import os
import sys

import openpyxl
from PySide6.QtCore import Qt
from PySide6.QtGui import QIcon, QAction
from PySide6.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QProgressBar, QFileDialog,
    QScrollArea, QLineEdit, QMessageBox, QTabWidget,
    QSystemTrayIcon, QStyle, QMenuBar,
)

from ui.url_row_widget import UrlRowWidget
from ui.config_panel import ConfigPanel
from ui.log_panel import LogPanel, QtLogHandler
from ui.theme import PALETTE, stylesheet
from scraper.worker import CrawlWorker


def _resource(rel: str) -> str:
    if getattr(sys, "frozen", False):
        base = sys._MEIPASS  # type: ignore[attr-defined]
    else:
        base = os.path.dirname(os.path.dirname(__file__))
    return os.path.join(base, rel)

logger = logging.getLogger(__name__)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("TikTok Image Scraper")
        self.setMinimumSize(1000, 700)
        self.resize(1280, 860)
        self._rows: list[UrlRowWidget] = []
        self._worker: CrawlWorker | None = None
        self._dest_folder = os.path.expanduser("~/Desktop/tiktok_images")

        icon_path = _resource(os.path.join("assets", "icon.ico"))
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))

        self.setStyleSheet(stylesheet())
        self._setup_logging()
        self._setup_tray()
        self._build_ui()
        self._build_menu()

    # ── Menu bar ────────────────────────────────────────────────────────────

    def _build_menu(self):
        menu_bar = self.menuBar()
        menu_bar.setStyleSheet(f"""
            QMenuBar {{ background: {PALETTE['bg_surface']}; color: {PALETTE['text']};
                        border-bottom: 1px solid {PALETTE['border']}; }}
            QMenuBar::item:selected {{ background: {PALETTE['bg_card']}; }}
            QMenu {{ background: {PALETTE['bg_surface']}; color: {PALETTE['text']};
                     border: 1px solid {PALETTE['border']}; }}
            QMenu::item:selected {{ background: {PALETTE['bg_card']}; }}
        """)
        help_menu = menu_bar.addMenu("Help")
        about_action = QAction("About", self)
        about_action.triggered.connect(self._on_about)
        help_menu.addAction(about_action)

    def _on_about(self):
        from ui.about_dialog import AboutDialog
        AboutDialog(self).exec()

    # ── Logging setup ───────────────────────────────────────────────────────

    def _setup_logging(self):
        self._log_handler = QtLogHandler()
        self._log_handler.setLevel(logging.DEBUG)
        logging.getLogger().addHandler(self._log_handler)

    # ── System tray ─────────────────────────────────────────────────────────

    def _setup_tray(self):
        icon = self.style().standardIcon(QStyle.StandardPixmap.SP_MessageBoxWarning)
        self._tray = QSystemTrayIcon(icon, self)
        self._tray.setToolTip("TikTok Image Scraper")
        if QSystemTrayIcon.isSystemTrayAvailable():
            self._tray.show()

    # ── UI Construction ──────────────────────────────────────────────────────

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        tabs = QTabWidget()
        tabs.addTab(self._build_scraper_tab(), "  Scraper  ")
        tabs.addTab(self._build_settings_tab(), "  Settings  ")
        root.addWidget(tabs)

        # Connect log handler to log panel
        self._log_handler.log_emitted.connect(self._log_panel.append_log)

    def _build_scraper_tab(self) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(12, 10, 12, 8)
        layout.setSpacing(8)

        layout.addLayout(self._build_toolbar())
        layout.addWidget(self._build_rows_area(), stretch=1)
        layout.addLayout(self._build_save_bar())

        self._log_panel = LogPanel()
        layout.addWidget(self._log_panel)

        return tab

    def _build_settings_tab(self) -> QWidget:
        self._config_panel = ConfigPanel()
        return self._config_panel

    def _build_toolbar(self) -> QHBoxLayout:
        bar = QHBoxLayout()
        bar.setSpacing(8)

        self._load_btn = QPushButton("Load URLs from file")
        self._load_btn.clicked.connect(self._on_load_file)
        bar.addWidget(self._load_btn)

        self._tmpl_btn = QPushButton("Download Template")
        self._tmpl_btn.clicked.connect(self._on_download_template)
        bar.addWidget(self._tmpl_btn)

        sep = QWidget()
        sep.setFixedWidth(1)
        sep.setStyleSheet(f"background: {PALETTE['border']};")
        bar.addWidget(sep)

        self._crawl_btn = QPushButton("Crawl All")
        self._crawl_btn.setObjectName("primary")
        self._crawl_btn.setEnabled(False)
        self._crawl_btn.clicked.connect(self._on_crawl)
        bar.addWidget(self._crawl_btn)

        self._stop_btn = QPushButton("Stop")
        self._stop_btn.setObjectName("danger")
        self._stop_btn.setEnabled(False)
        self._stop_btn.clicked.connect(self._on_stop)
        bar.addWidget(self._stop_btn)

        self._progress = QProgressBar()
        self._progress.setFixedHeight(16)
        self._progress.setMinimum(0)
        self._progress.setMaximum(1)
        self._progress.setValue(0)
        bar.addWidget(self._progress, stretch=1)

        self._progress_label = QLabel("0 / 0")
        self._progress_label.setFixedWidth(64)
        self._progress_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._progress_label.setStyleSheet(f"color: {PALETTE['text_dim']}; font-size: 11px;")
        bar.addWidget(self._progress_label)

        return bar

    def _build_rows_area(self) -> QScrollArea:
        self._rows_widget = QWidget()
        self._rows_layout = QVBoxLayout(self._rows_widget)
        self._rows_layout.setContentsMargins(0, 0, 0, 0)
        self._rows_layout.setSpacing(0)
        self._rows_layout.addStretch()

        scroll = QScrollArea()
        scroll.setWidget(self._rows_widget)
        scroll.setWidgetResizable(True)
        return scroll

    def _build_save_bar(self) -> QHBoxLayout:
        bar = QHBoxLayout()
        bar.setSpacing(8)

        sel_all = QPushButton("Select All")
        sel_all.clicked.connect(lambda: self._select_all(True))
        bar.addWidget(sel_all)

        desel_all = QPushButton("Deselect All")
        desel_all.clicked.connect(lambda: self._select_all(False))
        bar.addWidget(desel_all)

        bar.addStretch()

        dest_lbl = QLabel("Save to:")
        dest_lbl.setObjectName("section")
        bar.addWidget(dest_lbl)

        self._dest_edit = QLineEdit(self._dest_folder)
        self._dest_edit.setReadOnly(True)
        self._dest_edit.setMinimumWidth(280)
        bar.addWidget(self._dest_edit)

        browse_btn = QPushButton("Browse…")
        browse_btn.clicked.connect(self._on_browse_dest)
        bar.addWidget(browse_btn)

        self._save_btn = QPushButton("Save Selected (0)")
        self._save_btn.setObjectName("primary")
        self._save_btn.setEnabled(False)
        self._save_btn.clicked.connect(self._on_save)
        bar.addWidget(self._save_btn)

        return bar

    # ── Helpers ──────────────────────────────────────────────────────────────

    def _add_row(self, url: str, index: int, prefix: str = ""):
        row = UrlRowWidget(url, index, self._rows_widget)
        if prefix:
            row.prefix_edit.setText(prefix)
        self._rows_layout.insertWidget(self._rows_layout.count() - 1, row)
        self._rows.append(row)

    def _clear_rows(self):
        for row in self._rows:
            row.setParent(None)
        self._rows.clear()

    # ── Slots ────────────────────────────────────────────────────────────────

    def _on_load_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Open URL list", "",
            "Excel files (*.xlsx);;Text files (*.txt);;All files (*)"
        )
        if not path:
            return

        self._clear_rows()
        if path.lower().endswith(".xlsx"):
            entries = self._parse_xlsx(path)
        else:
            entries = self._parse_txt(path)

        for i, (url, prefix) in enumerate(entries, 1):
            self._add_row(url, i, prefix)

        self._crawl_btn.setEnabled(bool(entries))
        self._progress.setMaximum(max(len(entries), 1))
        self._progress.setValue(0)
        self._progress_label.setText(f"0 / {len(entries)}")
        logger.info("Loaded %d URL(s) from %s", len(entries), os.path.basename(path))

    def _parse_xlsx(self, path: str) -> list[tuple[str, str]]:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        url_col, prefix_col = 0, 1
        header = [str(c).lower().strip() if c is not None else "" for c in rows[0]]
        if "url" in header:
            url_col = header.index("url")
        if "prefix" in header:
            prefix_col = header.index("prefix")
        data_rows = rows[1:] if "url" in header else rows

        entries = []
        for row in data_rows:
            url = str(row[url_col]).strip() if url_col < len(row) and row[url_col] else ""
            prefix = str(row[prefix_col]).strip() if prefix_col < len(row) and row[prefix_col] else ""
            if url.startswith("http"):
                entries.append((url, prefix))
        return entries

    def _parse_txt(self, path: str) -> list[tuple[str, str]]:
        with open(path, encoding="utf-8") as f:
            lines = [l.strip() for l in f if l.strip()]
        entries = []
        for line in lines:
            parts = line.split(None, 1)
            url = parts[-1] if parts else ""
            if url.startswith("http"):
                entries.append((url, ""))
        return entries

    def _on_download_template(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Template", "tiktok_template.xlsx", "Excel files (*.xlsx)"
        )
        if not path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["url", "prefix"])
        ws.append(["https://vt.tiktok.com/EXAMPLE/", "product_name"])
        wb.save(path)
        QMessageBox.information(self, "Template Saved", f"Template saved to:\n{path}")

    def _on_crawl(self):
        if not self._rows:
            return
        self._crawl_btn.setEnabled(False)
        self._stop_btn.setEnabled(True)
        self._progress.setValue(0)
        self._progress.setMaximum(len(self._rows))

        for row in self._rows:
            row.set_status("Queued…")

        config = self._config_panel.get_config()
        logger.info(
            "Starting crawl: %d URL(s), human_sim=%s, delay=%.1f–%.1fs",
            len(self._rows), config["human_sim"], config["min_delay"], config["max_delay"],
        )

        urls = [row.url for row in self._rows]
        self._worker = CrawlWorker(urls, config=config)
        self._worker.images_ready.connect(self._on_images_ready)
        self._worker.status_update.connect(self._on_status_update)
        self._worker.progress.connect(self._on_progress)
        self._worker.captcha_detected.connect(self._on_captcha_detected)
        self._worker.finished.connect(self._on_crawl_finished)
        self._worker.start()

    def _on_stop(self):
        if self._worker:
            self._worker.stop()
        self._stop_btn.setEnabled(False)

    def _on_images_ready(self, url: str, pairs: list):
        for row in self._rows:
            if row.url == url:
                row.add_images(pairs)
                break
        self.update_save_button()

    def _on_status_update(self, url: str, text: str):
        for row in self._rows:
            if row.url == url:
                row.set_status(text)
                break

    def _on_progress(self, done: int, total: int):
        self._progress.setValue(done)
        self._progress_label.setText(f"{done} / {total}")

    def _on_captcha_detected(self, url: str):
        config = self._config_panel.get_config()

        if config.get("sound_notify"):
            try:
                import winsound
                winsound.MessageBeep(winsound.MB_ICONEXCLAMATION)
            except Exception:
                pass

        if config.get("tray_notify") and QSystemTrayIcon.isSystemTrayAvailable():
            short_url = url[:70] + ("…" if len(url) > 70 else "")
            self._tray.showMessage(
                "CAPTCHA Detected!",
                f"Please solve the security check in the browser window.\n{short_url}",
                QSystemTrayIcon.MessageIcon.Warning,
                10_000,
            )

        logger.warning("CAPTCHA detected — waiting for user to solve: %s", url)

    def _on_crawl_finished(self):
        self._crawl_btn.setEnabled(True)
        self._stop_btn.setEnabled(False)
        logger.info("All URLs processed")

    def _on_browse_dest(self):
        folder = QFileDialog.getExistingDirectory(self, "Select destination folder", self._dest_folder)
        if folder:
            self._dest_folder = folder
            self._dest_edit.setText(folder)

    def _select_all(self, state: bool):
        for row in self._rows:
            row.select_all(state)
        self.update_save_button()

    def update_save_button(self):
        total = sum(row.selected_count() for row in self._rows)
        self._save_btn.setText(f"Save Selected ({total})")
        self._save_btn.setEnabled(total > 0)

    def _on_save(self):
        dest = self._dest_edit.text().strip() or self._dest_folder
        os.makedirs(dest, exist_ok=True)

        saved = 0
        errors = 0
        for row in self._rows:
            selected = [t for t in row.thumbs if t.selected]
            if not selected:
                continue
            prefix = row.prefix_edit.text().strip() or "image"
            for i, thumb in enumerate(selected, 1):
                filename = f"{prefix}_{i:03d}.jpg"
                filepath = os.path.join(dest, filename)
                stem, ext = os.path.splitext(filepath)
                counter = 1
                while os.path.exists(filepath):
                    filepath = f"{stem}_{counter}{ext}"
                    counter += 1
                ok = thumb.full_pixmap.save(filepath, "JPEG", 92)
                if ok:
                    saved += 1
                else:
                    errors += 1

        logger.info("Saved %d image(s) to %s (%d error(s))", saved, dest, errors)
        msg = f"Saved {saved} image(s) to:\n{dest}"
        if errors:
            msg += f"\n\n⚠ {errors} image(s) could not be saved."
        QMessageBox.information(self, "Save Complete", msg)
